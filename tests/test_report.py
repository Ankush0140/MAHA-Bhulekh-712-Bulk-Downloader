import pytest
from pathlib import Path
import tempfile
import csv
import openpyxl
from fastapi.testclient import TestClient

from app.models.db import Base, Job, JobRecord
from app.models.enums import JobStatus, RecordStatus, ErrorCategory
from app.services.report import generate_job_reports
from app.config import BASE_DIR
from app.main import app
from app.dependencies import get_db

@pytest.fixture
def temp_db_session():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    db = TestingSessionLocal()
    yield db
    db.close()

@pytest.fixture
def temp_output_dir():
    with tempfile.TemporaryDirectory() as tmpdir:
        yield Path(tmpdir)

@pytest.fixture
def sample_job_and_records(temp_db_session):
    job = Job(
        division_text="Pune Division", division_value="pune",
        district_text="पुणे", district_value="25",
        taluka_text="हवेली", taluka_value="1",
        village_text="पुणे (म.न.पा.)", village_value="1",
        status=JobStatus.COMPLETED_WITH_ERRORS,
        total_records=3,
        completed_records=2,
        failed_records=1,
        empty_records=0
    )
    temp_db_session.add(job)
    temp_db_session.commit()
    
    # 1. Unicode record
    r1 = JobRecord(
        job_id=job.id,
        survey_identifier_original="30/2/अ",
        survey_identifier_safe="30_2_अ",
        status=RecordStatus.COMPLETED,
        output_relative_path="पुणे/हवेली/पुणे (म.न.पा.)/30_2_अ.pdf",
        attempt_count=1
    )
    
    # 2. Subdivision
    r2 = JobRecord(
        job_id=job.id,
        survey_identifier_original="59/2/1",
        survey_identifier_safe="59_2_1",
        status=RecordStatus.COMPLETED,
        output_relative_path="पुणे/हवेली/पुणे (म.न.पा.)/59_2_1.pdf",
        attempt_count=1
    )
    
    # 3. Formula Injection Attack
    r3 = JobRecord(
        job_id=job.id,
        survey_identifier_original="=cmd|' /C calc'!A0",
        survey_identifier_safe="cmd_C_calc_A0",
        status=RecordStatus.FAILED,
        last_error_code=ErrorCategory.PDF_WRITE_ERROR,
        last_error_message="+some error message",
        attempt_count=3
    )
    
    temp_db_session.add_all([r1, r2, r3])
    temp_db_session.commit()
    
    return job.id, temp_db_session

def test_csv_generation(sample_job_and_records, temp_output_dir):
    job_id, db = sample_job_and_records
    
    result = generate_job_reports(db, job_id, temp_output_dir)
    assert result is not None
    
    csv_rel = result["csv"]
    csv_full = temp_output_dir / csv_rel
    
    assert csv_full.exists()
    
    # check encoding BOM
    with open(csv_full, "rb") as f:
        content = f.read()
        assert content.startswith(b'\xef\xbb\xbf')
        
    with open(csv_full, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        
        assert len(rows) == 3
        
        # Unicode exact match
        assert rows[0]["Survey/Gat/Hissa"] == "30/2/अ"
        assert rows[0]["Status"] == "COMPLETED"
        assert rows[0]["Attempts"] == "1"
        assert "पुणे" in rows[0]["Output File"]
        
        # Subdivision exact match
        assert rows[1]["Survey/Gat/Hissa"] == "59/2/1"
        
        # Formula injection safe check
        assert rows[2]["Survey/Gat/Hissa"] == "'=cmd|' /C calc'!A0"
        assert rows[2]["Error Message"] == "'+some error message"
        
        # Sensitive columns check
        sensitive_terms = ["traceback", "stack_trace", "mobile", "captcha", "cookie", "token", "viewstate", "eventvalidation", "base64", "html"]
        headers_str = " ".join(rows[0].keys()).lower()
        for term in sensitive_terms:
            assert term not in headers_str, f"Found sensitive term {term} in CSV headers"

def test_excel_generation(sample_job_and_records, temp_output_dir):
    job_id, db = sample_job_and_records
    
    result = generate_job_reports(db, job_id, temp_output_dir)
    xlsx_rel = result["xlsx"]
    xlsx_full = temp_output_dir / xlsx_rel
    
    wb = openpyxl.load_workbook(xlsx_full)
    
    assert "Summary" in wb.sheetnames
    assert "Records" in wb.sheetnames
    
    ws_summary = wb["Summary"]
    summary_dict = {row[0]: row[1] for row in ws_summary.iter_rows(values_only=True)}
    assert summary_dict["Job ID"] == job_id
    assert summary_dict["District"] == "पुणे"
    assert summary_dict["Total Records"] == 3
    assert summary_dict["Failed"] == 1
    
    ws_records = wb["Records"]
    rows = list(ws_records.iter_rows(values_only=True))
    
    headers = rows[0]
    assert "Survey/Gat/Hissa" in headers
    
    headers_str = " ".join([str(h) for h in headers if h]).lower()
    sensitive_terms = ["traceback", "stack_trace", "mobile", "captcha", "cookie", "token", "viewstate", "eventvalidation", "base64", "html"]
    for term in sensitive_terms:
        assert term not in headers_str, f"Found sensitive term {term} in XLSX headers"
    
    data_rows = rows[1:]
    assert len(data_rows) == 3
    
    # record 1
    assert data_rows[0][1] == "30/2/अ"
    
    # formula injection record
    assert data_rows[2][1] == "'=cmd|' /C calc'!A0"
    
def test_regeneration_no_duplicate_rows(sample_job_and_records, temp_output_dir):
    job_id, db = sample_job_and_records
    
    res1 = generate_job_reports(db, job_id, temp_output_dir)
    csv_path = temp_output_dir / res1["csv"]
    
    with open(csv_path, "r", encoding="utf-8-sig") as f:
        assert len(list(csv.DictReader(f))) == 3
        
    # regenerate
    res2 = generate_job_reports(db, job_id, temp_output_dir)
    with open(csv_path, "r", encoding="utf-8-sig") as f:
        assert len(list(csv.DictReader(f))) == 3

from unittest.mock import patch

def test_atomic_replacement_csv_failure(sample_job_and_records, temp_output_dir):
    job_id, db = sample_job_and_records
    
    # 1. Generate valid reports
    res1 = generate_job_reports(db, job_id, temp_output_dir)
    csv_path = temp_output_dir / res1["csv"]
    
    with open(csv_path, "w", encoding="utf-8") as f:
        f.write("OLD VALID CSV")
        
    # 2. Simulate failure during CSV generation
    with patch("csv.writer", side_effect=Exception("Simulated CSV Write Error")):
        res2 = generate_job_reports(db, job_id, temp_output_dir)
        
    # 3. Check that the existing file is intact
    with open(csv_path, "r", encoding="utf-8") as f:
        assert f.read() == "OLD VALID CSV"
        
    # 4. Ensure no temp files are left in the dir
    dir_path = csv_path.parent
    temp_files = list(dir_path.glob("*.csv"))
    assert len(temp_files) == 1
    assert temp_files[0].name == csv_path.name

def test_atomic_replacement_xlsx_failure(sample_job_and_records, temp_output_dir):
    job_id, db = sample_job_and_records
    
    # 1. Generate valid reports
    res1 = generate_job_reports(db, job_id, temp_output_dir)
    xlsx_path = temp_output_dir / res1["xlsx"]
    
    with open(xlsx_path, "wb") as f:
        f.write(b"OLD VALID XLSX")
        
    # 2. Simulate failure during XLSX generation
    with patch("openpyxl.Workbook.save", side_effect=Exception("Simulated XLSX Write Error")):
        res2 = generate_job_reports(db, job_id, temp_output_dir)
        
    # 3. Check that the existing file is intact
    with open(xlsx_path, "rb") as f:
        assert f.read() == b"OLD VALID XLSX"
        
    # 4. Ensure no temp files are left in the dir
    dir_path = xlsx_path.parent
    temp_files = list(dir_path.glob("*.xlsx"))
    assert len(temp_files) == 1
    assert temp_files[0].name == xlsx_path.name

# API Tests
client = TestClient(app)

from sqlalchemy.pool import StaticPool
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

test_api_engine = create_engine(
    "sqlite:///:memory:", 
    connect_args={"check_same_thread": False},
    poolclass=StaticPool
)
TestApiSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_api_engine)

def override_get_db():
    db = TestApiSessionLocal()
    try:
        yield db
    finally:
        db.close()

def test_api_generate_and_download(temp_output_dir, monkeypatch):
    Base.metadata.create_all(bind=test_api_engine)
    app.dependency_overrides[get_db] = override_get_db
    try:
        monkeypatch.setattr("app.routes.jobs.BASE_DIR", temp_output_dir)
        
        # Create Job
        resp = client.post("/api/jobs", json={
            "division_text": "Pune", "division_value": "pune",
            "district_text": "Pune", "district_value": "25",
            "taluka_text": "Haveli", "taluka_value": "1",
            "village_text": "Pune", "village_value": "1"
        })
        job_id = resp.json()["id"]
        
        # Reports 404 before generated
        resp = client.get(f"/api/jobs/{job_id}/reports/csv")
        assert resp.status_code == 404
        
        resp = client.get(f"/api/jobs/{job_id}/reports/xlsx")
        assert resp.status_code == 404
        
        # Generate
        resp = client.post(f"/api/jobs/{job_id}/reports")
        assert resp.status_code == 200
        data = resp.json()
        assert "csv" in data
        assert "xlsx" in data
        
        # Download
        resp_csv = client.get(f"/api/jobs/{job_id}/reports/csv")
        assert resp_csv.status_code == 200
        assert resp_csv.headers["content-type"].startswith("text/csv")
        
        resp_xlsx = client.get(f"/api/jobs/{job_id}/reports/xlsx")
        assert resp_xlsx.status_code == 200
        assert "spreadsheetml" in resp_xlsx.headers["content-type"]
    finally:
        app.dependency_overrides.clear()
        
def test_api_invalid_job():
    app.dependency_overrides[get_db] = override_get_db
    try:
        resp = client.post("/api/jobs/999/reports")
        assert resp.status_code == 404
    finally:
        app.dependency_overrides.clear()
