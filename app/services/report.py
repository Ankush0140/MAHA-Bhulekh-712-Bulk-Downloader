import csv
import logging
from pathlib import Path
from typing import Dict, Any, Optional
import tempfile
import shutil
import openpyxl
from openpyxl.styles import Font, Alignment
from sqlalchemy.orm import Session

from app.models.db import Job, JobRecord
from app.services.db_service import get_job
from app.services.filename import sanitize_filename_component

logger = logging.getLogger(__name__)

def safe_string(value: Any) -> str:
    """Sanitize strings to prevent CSV/Excel formula injection."""
    if value is None:
        return ""
    v = str(value)
    if v.startswith(("=", "+", "-", "@")):
        return "'" + v
    return v

def _get_report_paths(job: Job, base_output_dir: Path) -> tuple[Path, Path, str, str]:
    dist = sanitize_filename_component(job.district_text)
    tal = sanitize_filename_component(job.taluka_text)
    vil = sanitize_filename_component(job.village_text)
    
    dir_path = base_output_dir / dist / tal / vil
    dir_path.mkdir(parents=True, exist_ok=True)
    
    prefix = f"{dist}_{tal}_{vil}_summary"
    csv_path = dir_path / f"{prefix}.csv"
    xlsx_path = dir_path / f"{prefix}.xlsx"
    
    rel_csv = csv_path.relative_to(base_output_dir).as_posix()
    rel_xlsx = xlsx_path.relative_to(base_output_dir).as_posix()
    
    return csv_path, xlsx_path, rel_csv, rel_xlsx

def generate_job_reports(db: Session, job_id: int, base_output_dir: Path) -> Optional[Dict[str, str]]:
    import os
    job = get_job(db, job_id)
    if not job:
        return None
        
    records = db.query(JobRecord).filter(JobRecord.job_id == job_id).order_by(JobRecord.id.asc()).all()
    
    csv_path, xlsx_path, rel_csv, rel_xlsx = _get_report_paths(job, base_output_dir)
    dir_path = csv_path.parent
    
    csv_success = False
    temp_csv_path = None
    try:
        fd, temp_csv_path_str = tempfile.mkstemp(dir=dir_path, suffix=".csv")
        temp_csv_path = Path(temp_csv_path_str)
        with os.fdopen(fd, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            headers = [
                "Survey/Gat/Hissa", "Status", "Attempts", "Output File", 
                "Error Category", "Error Message", 
                "Automated Processing Seconds", "Human Verification Seconds",
                "Division", "District", "Taluka", "Village"
            ]
            writer.writerow(headers)
            
            for r in records:
                writer.writerow([
                    safe_string(r.survey_identifier_original),
                    safe_string(r.status.value),
                    r.attempt_count,
                    safe_string(r.output_relative_path),
                    safe_string(r.last_error_code.value if r.last_error_code else ""),
                    safe_string(r.last_error_message),
                    round(r.automated_duration_seconds or 0, 2),
                    round(r.human_wait_seconds or 0, 2),
                    safe_string(job.division_text),
                    safe_string(job.district_text),
                    safe_string(job.taluka_text),
                    safe_string(job.village_text)
                ])
        
        os.replace(temp_csv_path, csv_path)
        csv_success = True
    except Exception as e:
        logger.error(f"Failed to generate CSV for job {job_id}: {e}")
        if temp_csv_path and temp_csv_path.exists():
            try:
                os.remove(temp_csv_path)
            except OSError:
                pass
                
    xlsx_success = False
    temp_xlsx_path = None
    try:
        fd, temp_xlsx_path_str = tempfile.mkstemp(dir=dir_path, suffix=".xlsx")
        temp_xlsx_path = Path(temp_xlsx_path_str)
        os.close(fd) # openpyxl requires a filename to save, so close fd
        
        wb = openpyxl.Workbook()
        
        # Summary Sheet
        ws_summary = wb.active
        ws_summary.title = "Summary"
        
        summary_data = [
            ("Job ID", job.id),
            ("Job Status", safe_string(job.status.value)),
            ("Division", safe_string(job.division_text)),
            ("District", safe_string(job.district_text)),
            ("Taluka", safe_string(job.taluka_text)),
            ("Village", safe_string(job.village_text)),
            ("Total Records", job.total_records),
            ("Completed", job.completed_records),
            ("Failed", job.failed_records),
            ("Empty", job.empty_records),
            ("Remaining", max(0, job.total_records - (job.completed_records + job.failed_records + job.empty_records))),
            ("Automated Processing Seconds", round(job.automated_processing_seconds or 0, 2)),
            ("Human Verification Seconds", round(job.human_wait_seconds or 0, 2)),
            ("Created At", str(job.created_at) if job.created_at else ""),
            ("Started At", str(job.started_at) if job.started_at else ""),
            ("Completed At", str(job.completed_at) if job.completed_at else ""),
            ("Numeric Discovery Support", "Supported"),
            ("Akshar Enumeration", "Unverified / Unsupported"),
            ("CAPTCHA", "Human-in-the-loop")
        ]
        
        for k, v in summary_data:
            ws_summary.append([k, v])
            
        # Bold first column
        for row in ws_summary.iter_rows(min_col=1, max_col=1):
            for cell in row:
                cell.font = Font(bold=True)
        
        ws_summary.column_dimensions['A'].width = 30
        ws_summary.column_dimensions['B'].width = 40
        
        # Records Sheet
        ws_records = wb.create_sheet(title="Records")
        record_headers = [
            "ID", "Survey/Gat/Hissa", "Status", "Attempts", "Output File",
            "Error Category", "Error Message", "Automated Processing Seconds", "Human Verification Seconds"
        ]
        ws_records.append(record_headers)
        
        # Format header
        for cell in ws_records[1]:
            cell.font = Font(bold=True)
        ws_records.freeze_panes = "A2"
        
        for r in records:
            ws_records.append([
                r.id,
                safe_string(r.survey_identifier_original),
                safe_string(r.status.value),
                r.attempt_count,
                safe_string(r.output_relative_path),
                safe_string(r.last_error_code.value if r.last_error_code else ""),
                safe_string(r.last_error_message),
                round(r.automated_duration_seconds or 0, 2),
                round(r.human_wait_seconds or 0, 2)
            ])
            
        ws_records.auto_filter.ref = f"A1:I{len(records)+1}"
            
        for col in ['A', 'C', 'D', 'E', 'F', 'H', 'I']:
            ws_records.column_dimensions[col].width = 15
        ws_records.column_dimensions['B'].width = 25
        ws_records.column_dimensions['G'].width = 40
        
        wb.save(temp_xlsx_path)
        os.replace(temp_xlsx_path, xlsx_path)
        xlsx_success = True
    except Exception as e:
        logger.error(f"Failed to generate Excel for job {job_id}: {e}")
        if temp_xlsx_path and temp_xlsx_path.exists():
            try:
                os.remove(temp_xlsx_path)
            except OSError:
                pass
                
    if csv_success or xlsx_success:
        return {
            "csv": rel_csv if csv_success else "",
            "xlsx": rel_xlsx if xlsx_success else ""
        }
    return None
